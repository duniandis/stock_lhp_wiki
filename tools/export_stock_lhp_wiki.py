import csv, os, sys
from datetime import date, datetime
from openpyxl import load_workbook

SOURCE = os.getenv("SOURCE_XLSX", "INPUT ANGKUTAN DOKUMENT 2026.xlsx")
OUTPUT = "stock_lhp_wiki.csv"
SHEET = "RECAP POSISI"

def clean(v):
    return "" if v is None else str(v).strip()

def n(v):
    if v in (None, ""): return 0
    return v

def fmt_num(v):
    try: return f"{float(n(v)):.4f}"
    except (TypeError, ValueError): return clean(v)

def fmt_date(v):
    if isinstance(v, (datetime, date)): return v.strftime("%d-%m-%Y")
    return clean(v)

def jenis_row(tipe, vals):
    # B:J => jenis, 40-49 btg/vol, 50-59 btg/vol, 60+ btg/vol, total btg/vol
    return [tipe, clean(vals[0]), *[fmt_num(x) for x in vals[1:]], "", "", "", "", "", ""]

def main():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"Sheet {SHEET!r} tidak ditemukan")
    ws = wb[SHEET]

    rows = []
    # STOCK LHP: B55:J65
    for r in range(55, 66):
        vals = [ws.cell(r, c).value for c in range(2, 11)]
        if not any(v not in (None, "") for v in vals): continue
        jenis = clean(vals[0])
        if r == 65: jenis = "TOTAL"
        vals[0] = jenis
        rows.append(jenis_row("STOCK", vals))

    # PEMILIRAN: B31:J41
    for r in range(31, 42):
        vals = [ws.cell(r, c).value for c in range(2, 11)]
        if not any(v not in (None, "") for v in vals): continue
        jenis = clean(vals[0])
        if r == 41: jenis = "TOTAL PEMILIRAN"
        vals[0] = jenis
        rows.append(jenis_row("PEMILIRAN", vals))

    # DETAIL: A70:F sampai baris terakhir yang berisi data.
    # Header Excel ada di A69:F69.
    blank_run = 0
    for r in range(70, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if not any(v not in (None, "") for v in vals):
            blank_run += 1
            if blank_run >= 20: break
            continue
        blank_run = 0
        no, no_seri, tanggal, btg, volume, tujuan = vals
        # Hanya baris dokumen yang memiliki nomor seri.
        if not clean(no_seri): continue
        rows.append([
            "DETAIL","","","","","","","","","",
            clean(no), clean(no_seri), fmt_date(tanggal), fmt_num(btg), fmt_num(volume), clean(tujuan)
        ])

    header = [
        "tipe","jenis","btg_40_49","vol_40_49","btg_50_59","vol_50_59",
        "btg_60_up","vol_60_up","total_btg","total_vol",
        "no","no_seri","tanggal","btg","volume","tujuan"
    ]
    tmp = OUTPUT + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    os.replace(tmp, OUTPUT)
    print(f"{OUTPUT}: {len(rows)} baris diekspor")

if __name__ == "__main__":
    try: main()
    except Exception as e:
        print("ERROR:", e, file=sys.stderr)
        sys.exit(1)
